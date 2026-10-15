#!/usr/bin/env python3
"""Stammdaten einer Energiegemeinschaft aus CSV oder Excel importieren.

Ersetzt im Labor das OctoMesh-Excel-Add-in (Befehl "Query Publish").

Kunden und Anlagen gehen ueber dieselbe GraphQL-Mutation wie das Add-in:

    runtime { runtimeQuery(rtId: <Abfrage>) { create(entities: [{ckTypeId, cells}]) } }
    runtime { runtimeQuery(rtId: <Abfrage>) { update(entities: [{rtId, item}]) } }

Die Abfragen saet EnergyCommunity.Base (seed-data/queries/master-data.yaml):
_energy_customers und _energy_operating_facilities. Jede Spalte der Datei
ist ein attributePath der Abfrage, z. B. `contact.lastName` oder
`parent.basicCity->zipcode`. Pfade mit `->` sind Verweise: der Server sucht
das Zielobjekt ueber den Wert (PLZ, Kundennummer).

Zaehlpunkte (Verbraucher, Erzeuger) gehen NICHT ueber die Abfragen
_energy_meter_consumer / _energy_meter_producer, weil die auf OctoMesh
3.4.141 unbrauchbar sind (Spaltenexpansion ASSET1006 wegen des
Zwei-Schritt-Verweises auf die Kundennummer, ausserdem fehlt das
Pflichtfeld carrierType). Das Skript nimmt dieselben Spalten, sucht die
Anlage selbst (Name + Kundennummer) und schreibt ueber die Entity-Mutation
runtime.energyCommunityConsumers / energyCommunityProducers.

Regeln wie beim Add-in:
  * Spalte `rtId` leer  -> Objekt wird neu angelegt.
  * Spalte `rtId` gefuellt -> Objekt wird aktualisiert (alle Spalten der Zeile).
  * Leere Zelle = Wert loeschen. Ausnahme: eine leere Verweis-Spalte (`->`)
    wird nicht mitgeschickt, die Zuordnung bleibt dann, wie sie ist.
Zusaetzlich zum Add-in:
  * Nach dem Import schreibt das Skript die rtIds in die Datei zurueck.
  * Fehlt die rtId, sucht das Skript ein vorhandenes Objekt ueber den
    fachlichen Schluessel (Kundennummer, Anlage = Name + Kundennummer,
    Zaehlpunktnummer). Ein zweiter Lauf legt also nichts doppelt an.
    Abschalten mit --ohne-schluesselabgleich.
  * Unveraenderte Zeilen werden nicht geschickt.
  * Verweise (Kundennummer, PLZ, Anlage) werden vorab geprueft, auch im
    Probelauf.
  * Vergibt ein Zaehler beim Anlegen einen anderen Wert (Kundennummer),
    traegt das Skript den Wert aus der Datei sofort nach.

Beispiele:
  scripts/import-stammdaten.py -c fh-cloud-lab_eg-urstein alle templates/stammdaten/ --probelauf
  scripts/import-stammdaten.py -c fh-cloud-lab_eg-urstein alle templates/stammdaten/
  scripts/import-stammdaten.py -c fh-cloud-lab_eg-urstein kunden meine-kunden.csv
  scripts/import-stammdaten.py -c fh-cloud-lab_eg-urstein alle stammdaten.xlsx
  scripts/import-stammdaten.py -c fh-cloud-lab_eg-urstein spalten

Voraussetzung: octo-cli mit einem angemeldeten Kontext auf den Tenant der
Gemeinschaft. Das Skript holt das Zugriffstoken aus diesem Kontext und gibt
es nie aus.
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import os
import re
import shutil
import ssl
import subprocess
import sys
import urllib.error
import urllib.request
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

try:
    from zoneinfo import ZoneInfo
    LOCAL_TZ = ZoneInfo("Europe/Vienna")
except Exception:  # pragma: no cover
    LOCAL_TZ = None

# --------------------------------------------------------------------------
# Typen. Die Reihenfolge ist die Import-Reihenfolge (Verweise zeigen nach oben).
# --------------------------------------------------------------------------

FAC_NAME = "parent.basicEnergyOperatingFacility->name"
FAC_CUST = "parent.basicEnergyOperatingFacility->customer.energyCommunityCustomer->customerNumber"
FAC_CK = "Basic.Energy/OperatingFacility"
CUST_REF = "customer.energyCommunityCustomer->customerNumber"
CITY_REF = "parent.basicCity->zipcode"

# CK-Enums in Schluesselreihenfolge (Basic 2.4, Basic.Energy 1.9, EnergyCommunity 4.8).
# In der Datei geht der Name (z. B. Solar) oder der Schluessel (1).
ENUMS: dict[str, list[str]] = {
    "contact.legalEntityType": ["NaturalPerson", "LegalPerson", "Company", "LocalAuthority"],
    "contact.salutation": ["Unknown", "Male", "Female", "NonBinary"],
    "taxProcedureCreditNote": ["SmallBusinessOwners", "ReverseCharge", "NoTaxProcedure", "FarmerTaxProcedure"],
    "facilityType": ["Unknown", "Household", "Business", "Industry", "PublicBuilding", "Storage"],
    "productionType": ["Unknown", "Solar", "HEP", "Wind", "Other", "Biomass", "CHP"],
    "carrierType": ["Unknown", "Electricity", "Gas"],
}


@dataclass(frozen=True)
class EntitySpec:
    """Zaehlpunkte: Lesen und Schreiben ueber die typisierte GraphQL-API."""
    ck_type: str
    query_field: str
    mutation_field: str
    columns: tuple[tuple[str, str], ...]   # (attributePath, Werttyp)
    value_field: str


@dataclass(frozen=True)
class ImportType:
    key: str
    query_name: str
    fallback_rtid: str
    natural_key: tuple[str, ...]
    sheet_names: tuple[str, ...]
    label: str
    entity: EntitySpec | None = None


METER_COMMON = (("meteringPointNumber", "STRING"), ("name", "STRING"),
                (FAC_NAME, "STRING"), (FAC_CUST, "STRING"),
                ("carrierType", "ENUM"), ("partitionFactor", "INT"))

TYPES: list[ImportType] = [
    ImportType("kunden", "_energy_customers", "682f0758900aec7bc66dcb74",
               ("customerNumber",), ("kunden", "customers"), "Kunden"),
    ImportType("anlagen", "_energy_operating_facilities", "682f0d3cce5a9c50f9241cc8",
               ("name", CUST_REF), ("anlagen", "anlage", "facilities"), "Anlagen"),
    ImportType("verbraucher", "_energy_meter_consumer", "682f1209ce5a9c50f9242880",
               ("meteringPointNumber",),
               ("verbraucher", "zaehlpunkte - verbraucher", "zählpunkte - verbraucher", "consumer"),
               "Zaehlpunkte Verbraucher",
               EntitySpec("EnergyCommunity/Consumer", "energyCommunityConsumer", "energyCommunityConsumers",
                          METER_COMMON[:2] + (("energyConsumption", "DOUBLE"),) + METER_COMMON[2:],
                          "energyConsumption")),
    ImportType("erzeuger", "_energy_meter_producer", "682f1218ce5a9c50f9242881",
               ("meteringPointNumber",),
               ("erzeuger", "zaehlpunkte - erzeuger", "zählpunkte - erzeuger", "producer"),
               "Zaehlpunkte Erzeuger",
               EntitySpec("EnergyCommunity/Producer", "energyCommunityProducer", "energyCommunityProducers",
                          METER_COMMON[:2] + (("energyProductionCapacity", "DOUBLE"), ("productionType", "ENUM"))
                          + METER_COMMON[2:],
                          "energyProductionCapacity productionType")),
]
TYPES_BY_KEY = {t.key: t for t in TYPES}

SPECIAL_COLUMNS = ("rtId", "ckTypeId", "rtVersion")
BATCH_SIZE = 500
PAGE_SIZE = 500


class UserError(Exception):
    """Fehler, der ohne Stacktrace gemeldet wird."""


def fail(msg: str) -> None:
    raise UserError(msg)


# --------------------------------------------------------------------------
# octo-cli-Kontext und Token
# --------------------------------------------------------------------------


@dataclass
class Connection:
    graphql_url: str
    tenant: str
    token: str = field(repr=False)
    ssl_context: ssl.SSLContext | None = field(default=None, repr=False)


def contexts_file() -> Path:
    home = (os.environ.get("OCTO_CLI_HOME") or "").strip()
    return (Path(home) if home else Path.home()) / ".octo-cli" / "contexts.json"


def refresh_token(context: str) -> None:
    """`octo-cli --context <ctx> -c AuthStatus` erneuert ein abgelaufenes Token still.
    Die Ausgabe wird verworfen: AuthStatus druckt das Token."""
    if shutil.which("octo-cli") is None:
        fail("octo-cli nicht gefunden. Bitte installieren und anmelden (docs/06-stammdaten-import.md).")
    try:
        result = subprocess.run(["octo-cli", "--context", context, "-c", "AuthStatus"],
                                stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                                stdin=subprocess.DEVNULL, timeout=120)
    except subprocess.TimeoutExpired:
        fail("octo-cli AuthStatus antwortet nicht (Timeout 120 s).")
    if result.returncode != 0:
        fail(f"Anmeldung fuer Kontext '{context}' ungueltig. Bitte neu anmelden:\n"
             f"  octo-cli --context {context} -c LogIn -i --if-needed")


def load_connection(args: argparse.Namespace) -> Connection:
    if not args.kein_refresh:
        refresh_token(args.context)
    path = contexts_file()
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError:
        fail(f"Keine octo-cli-Kontexte gefunden ({path}). Erst `octo-cli -c AddContext ...` und LogIn.")
    except json.JSONDecodeError:
        fail(f"{path} ist kein gueltiges JSON.")
    contexts = {k.lower(): v for k, v in (data.get("Contexts") or {}).items()}
    entry = contexts.get(args.context.lower())
    if entry is None:
        fail(f"Kontext '{args.context}' gibt es nicht. Vorhandene: octo-cli -c ListContexts")
    opts = entry.get("OctoToolOptions") or {}
    auth = entry.get("Authentication") or {}
    token = auth.get("AccessToken")
    if not token:
        fail(f"Kontext '{args.context}' hat kein Token. Bitte anmelden: octo-cli --context {args.context} -c LogIn -i")
    expires = auth.get("AccessTokenExpiresAt")
    if expires:
        try:
            exp = dt.datetime.fromisoformat(re.sub(r"(\.\d{6})\d+", r"\1", expires).replace("Z", "+00:00"))
            if exp.tzinfo is None:
                exp = exp.replace(tzinfo=dt.timezone.utc)
            if exp < dt.datetime.now(dt.timezone.utc):
                fail(f"Token von '{args.context}' ist abgelaufen und liess sich nicht erneuern. Bitte anmelden:\n"
                     f"  octo-cli --context {args.context} -c LogIn -i")
        except ValueError:
            pass
    assets = args.assets_url or opts.get("AssetServiceUrl")
    tenant = args.tenant or opts.get("TenantId")
    if not assets or not tenant:
        fail(f"Kontext '{args.context}' nennt keinen Asset-Dienst oder Tenant. --assets-url / --tenant angeben.")
    if not assets.endswith("/"):
        assets += "/"
    ctx = None
    if args.unsicher:
        ctx = ssl.create_default_context()
        ctx.check_hostname = False
        ctx.verify_mode = ssl.CERT_NONE
    return Connection(graphql_url=f"{assets}tenants/{tenant}/GraphQL", tenant=tenant, token=token, ssl_context=ctx)


# --------------------------------------------------------------------------
# GraphQL
# --------------------------------------------------------------------------

Q_FIND_QUERY = """
query findQuery($name: SimpleScalar) {
  runtime {
    systemSimpleRtQuery(fieldFilter: [{attributePath: "name", operator: EQUALS, comparisonValue: $name}]) {
      items { rtId name queryCkTypeId }
    }
  }
}"""

# navigationFilterMode INCLUDE: ohne ihn blendet der Server Zeilen mit Verweis-Spalten aus
# (lokal beobachtet: frisch angelegte Anlagen fehlten trotz korrekter Zuordnung).
Q_RUN_QUERY = """
query runQuery($queryRtId: OctoObjectId!, $after: String, $first: Int) {
  runtime {
    runtimeQuery(rtId: $queryRtId) {
      items {
        queryRtId
        associatedCkTypeId
        columns { attributePath attributeValueType }
        rows(after: $after, first: $first, navigationFilterMode: INCLUDE) {
          items {
            ... on RtSimpleQueryRow { rtId rtVersion }
            ckTypeId
            cells { items { attributePath value } }
          }
          pageInfo { hasNextPage endCursor }
          totalCount
        }
      }
    }
  }
}"""

M_CREATE = """
mutation createQueryData($queryRtId: OctoObjectId!, $entities: [RtQueryRowInput!]!) {
  runtime {
    runtimeQuery(rtId: $queryRtId) {
      create(entities: $entities) {
        ... on RtSimpleQueryRow { rtId rtVersion }
        ckTypeId
        cells { items { attributePath value } }
      }
    }
  }
}"""

M_UPDATE = """
mutation updateQueryData($queryRtId: OctoObjectId!, $entities: [RtQueryRowUpdate!]!) {
  runtime {
    runtimeQuery(rtId: $queryRtId) {
      update(entities: $entities) {
        ... on RtSimpleQueryRow { rtId rtVersion }
        ckTypeId
      }
    }
  }
}"""

Q_CITIES = """
query cities($after: String, $first: Int) {
  runtime { basicCity(after: $after, first: $first) {
    items { zipcode } pageInfo { hasNextPage endCursor } } }
}"""


def q_entities(spec: EntitySpec) -> str:
    return f"""
query meters($after: String, $first: Int) {{
  runtime {{ {spec.query_field}(after: $after, first: $first) {{
    pageInfo {{ hasNextPage endCursor }}
    items {{ rtId meteringPointNumber name carrierType partitionFactor {spec.value_field}
      parent(ckTypeIds: ["{FAC_CK}"]) {{ items {{ ... on BasicEnergyOperatingFacility {{ rtId name
        customer(ckTypeIds: ["EnergyCommunity/Customer"]) {{ items {{ ... on EnergyCommunityCustomer {{ customerNumber }} }} }}
      }} }} }}
    }} }} }}
}}"""


def m_entities(spec: EntitySpec, op: str) -> str:
    input_type = spec.ck_type.replace("/", "") + ("Input" if op == "create" else "InputUpdate")
    return f"""
mutation meters($entities: [{input_type}]!) {{
  runtime {{ {spec.mutation_field} {{ {op}(entities: $entities) {{ rtId meteringPointNumber }} }} }}
}}"""


class GraphQlError(UserError):
    def __init__(self, msg: str, errors: list[dict[str, Any]]):
        super().__init__(msg)
        self.errors = errors


def graphql(conn: Connection, query: str, variables: dict[str, Any]) -> dict[str, Any]:
    body = json.dumps({"query": query, "variables": variables}).encode("utf-8")
    req = urllib.request.Request(conn.graphql_url, data=body, method="POST", headers={
        "Content-Type": "application/json",
        "Accept": "application/json",
        "Authorization": f"Bearer {conn.token}",
    })
    try:
        with urllib.request.urlopen(req, timeout=300, context=conn.ssl_context) as resp:
            raw = resp.read()
    except urllib.error.HTTPError as e:
        hint = {401: "Token ungueltig oder fuer einen anderen Tenant ausgestellt.",
                403: "Keine Berechtigung auf diesem Tenant.",
                404: "Tenant oder GraphQL-Endpunkt nicht gefunden (Tenant-Id im Kontext pruefen)."}.get(e.code, "")
        fail(f"HTTP {e.code} von {conn.graphql_url}. {hint}".strip())
    except urllib.error.URLError as e:
        reason = getattr(e, "reason", e)
        extra = " (eigenes Zertifikat? --unsicher nur lokal verwenden)" if isinstance(reason, ssl.SSLError) else ""
        fail(f"Keine Verbindung zu {conn.graphql_url}: {reason}{extra}")
    try:
        doc = json.loads(raw)
    except json.JSONDecodeError:
        fail(f"Antwort von {conn.graphql_url} ist kein JSON (Proxy-Fehlerseite?).")
    # GraphQL meldet Fehler mit HTTP 200 und einem errors-Feld - immer auswerten.
    errors = doc.get("errors") or []
    if errors:
        lines = []
        for err in errors:
            ext = err.get("extensions") or {}
            line = f"  - {(err.get('message') or str(err)).strip()}"
            if ext.get("code"):
                line += f" [{ext['code']}]"
            # OctoDetails nennt die eigentliche Ursache (Mapping-Fehler, Pflichtfeld ...).
            # Das Feld `details` ist ein .NET-Stacktrace und wird nicht ausgegeben.
            for d in ext.get("OctoDetails") or []:
                msg = d.get("message") if isinstance(d, dict) else str(d)
                det = d.get("details") if isinstance(d, dict) else None
                line += f"\n      {msg}"
                if det:
                    line += " " + ", ".join(str(x) for x in (det if isinstance(det, list) else [det]))
            lines.append(line)
        raise GraphQlError("Der Server hat die Anfrage abgelehnt:\n" + "\n".join(lines), errors)
    return doc.get("data") or {}


def paged(conn: Connection, query: str, variables: dict[str, Any], getter) -> list[dict[str, Any]]:
    out, after = [], None
    while True:
        conn_obj = getter(graphql(conn, query, {**variables, "first": PAGE_SIZE, "after": after})) or {}
        out.extend(i for i in conn_obj.get("items") or [] if i)
        info = conn_obj.get("pageInfo") or {}
        if not info.get("hasNextPage"):
            return out
        after = info.get("endCursor")


# --------------------------------------------------------------------------
# Bestand im Tenant lesen
# --------------------------------------------------------------------------


@dataclass
class Existing:
    query_rtid: str | None
    ck_type_id: str
    columns: dict[str, str]                 # attributePath -> Werttyp
    column_order: list[str]
    rows: dict[str, dict[str, Any]]         # rtId -> {attributePath: Wert}
    facility_of: dict[str, str] = field(default_factory=dict)   # Zaehlpunkt-rtId -> Anlagen-rtId


def resolve_query(conn: Connection, itype: ImportType) -> str:
    try:
        data = graphql(conn, Q_FIND_QUERY, {"name": itype.query_name})
        items = (((data.get("runtime") or {}).get("systemSimpleRtQuery") or {}).get("items")) or []
        items = [i for i in items if i and i.get("name") == itype.query_name]
        if len(items) == 1:
            return items[0]["rtId"]
        if len(items) > 1:
            fail(f"Abfrage '{itype.query_name}' gibt es {len(items)}-mal im Tenant. Bitte bereinigen.")
    except GraphQlError:
        pass
    return itype.fallback_rtid


def load_query(conn: Connection, itype: ImportType) -> Existing:
    rtid = resolve_query(conn, itype)
    rows: dict[str, dict[str, Any]] = {}
    columns: dict[str, str] = {}
    order: list[str] = []
    ck_type, after = "", None
    while True:
        data = graphql(conn, Q_RUN_QUERY, {"queryRtId": rtid, "first": PAGE_SIZE, "after": after})
        items = (((data.get("runtime") or {}).get("runtimeQuery") or {}).get("items")) or []
        if not items:
            fail(f"Abfrage '{itype.query_name}' ({rtid}) fehlt im Tenant {conn.tenant}. "
                 f"Ist EnergyCommunity.Base installiert?")
        q = items[0]
        ck_type = q.get("associatedCkTypeId") or ck_type
        if not columns:
            for c in q.get("columns") or []:
                columns[c["attributePath"]] = c.get("attributeValueType") or "STRING"
                order.append(c["attributePath"])
        page = q.get("rows") or {}
        for r in page.get("items") or []:
            rows[r["rtId"]] = {c["attributePath"]: c.get("value")
                               for c in ((r.get("cells") or {}).get("items") or [])}
        info = page.get("pageInfo") or {}
        if not info.get("hasNextPage"):
            break
        after = info.get("endCursor")
    return Existing(query_rtid=rtid, ck_type_id=ck_type, columns=columns, column_order=order, rows=rows)


def first_item(conn_obj: Any) -> dict[str, Any] | None:
    items = (conn_obj or {}).get("items") or []
    return items[0] if items else None


def load_entities(conn: Connection, spec: EntitySpec) -> Existing:
    items = paged(conn, q_entities(spec), {}, lambda d: (d.get("runtime") or {}).get(spec.query_field))
    rows: dict[str, dict[str, Any]] = {}
    facility_of: dict[str, str] = {}
    for it in items:
        values = {path: it.get(path) for path, _ in spec.columns if "->" not in path}
        fac = first_item(it.get("parent"))
        if fac:
            facility_of[it["rtId"]] = fac["rtId"]
            values[FAC_NAME] = fac.get("name")
            cust = first_item(fac.get("customer"))
            values[FAC_CUST] = cust.get("customerNumber") if cust else None
        rows[it["rtId"]] = values
    return Existing(query_rtid=None, ck_type_id=spec.ck_type, columns=dict(spec.columns),
                    column_order=[p for p, _ in spec.columns], rows=rows, facility_of=facility_of)


def load_existing(conn: Connection, itype: ImportType) -> Existing:
    return load_entities(conn, itype.entity) if itype.entity else load_query(conn, itype)


# --------------------------------------------------------------------------
# Verweise vorab pruefen (Kunden, PLZ, Anlagen)
# --------------------------------------------------------------------------


class References:
    """Bestand im Tenant plus das, was dieser Lauf neu anlegt (wichtig im Probelauf)."""

    def __init__(self, conn: Connection):
        self.conn = conn
        self._customers: set | None = None
        self._zipcodes: dict | None = None
        self._facilities: dict | None = None
        self.pending_customers: set = set()
        self.pending_facilities: set = set()

    def invalidate(self) -> None:
        self._customers = self._facilities = None

    def customers(self) -> set:
        if self._customers is None:
            q = load_query(self.conn, TYPES_BY_KEY["kunden"])
            self._customers = {norm(r.get("customerNumber")) for r in q.rows.values()}
        return self._customers

    def zipcodes(self) -> dict:
        if self._zipcodes is None:
            counts: dict = {}
            try:
                for c in paged(self.conn, Q_CITIES, {}, lambda d: (d.get("runtime") or {}).get("basicCity")):
                    k = norm(c.get("zipcode"))
                    counts[k] = counts.get(k, 0) + 1
            except GraphQlError:
                counts = {}
            self._zipcodes = counts
        return self._zipcodes

    def facilities(self) -> dict:
        if self._facilities is None:
            q = load_query(self.conn, TYPES_BY_KEY["anlagen"])
            idx: dict = {}
            for rtid, r in q.rows.items():
                idx.setdefault((norm(r.get("name")), norm(r.get(CUST_REF))), []).append(rtid)
            self._facilities = idx
        return self._facilities


# --------------------------------------------------------------------------
# Dateien lesen und schreiben
# --------------------------------------------------------------------------


@dataclass
class Table:
    path: Path
    sheet: str | None
    header: list[str]
    rows: list[list[Any]]
    first_line: int
    delimiter: str = ";"


def read_csv(path: Path) -> Table:
    text = path.read_text(encoding="utf-8-sig")
    sample = "\n".join(text.splitlines()[:5])
    delimiter = ";"
    if sample:
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=";,\t").delimiter
        except csv.Error:
            delimiter = ";" if sample.count(";") >= sample.count(",") else ","
    reader = list(csv.reader(text.splitlines(), delimiter=delimiter))
    if not reader:
        fail(f"{path}: Datei ist leer.")
    header = [h.strip() for h in reader[0]]
    return Table(path=path, sheet=None, header=header, rows=reader[1:], first_line=2, delimiter=delimiter)


def require_openpyxl():
    try:
        import openpyxl
        return openpyxl
    except ImportError:
        fail("Fuer .xlsx wird openpyxl gebraucht: python3 -m pip install openpyxl  (oder als CSV speichern)")


def read_xlsx(path: Path, sheet: str) -> Table:
    openpyxl = require_openpyxl()
    ws = openpyxl.load_workbook(path, data_only=True)[sheet]
    values = [list(r) for r in ws.iter_rows(values_only=True)]
    if not values:
        fail(f"{path} [{sheet}]: Blatt ist leer.")
    header = [("" if h is None else str(h).strip()) for h in values[0]]
    while header and header[-1] == "":
        header.pop()
    return Table(path=path, sheet=sheet, header=header, rows=[r[:len(header)] for r in values[1:]], first_line=2)


def xlsx_sheet_for(path: Path, itype: ImportType) -> str | None:
    names = {n.lower().strip(): n for n in require_openpyxl().load_workbook(path, read_only=True).sheetnames}
    for cand in itype.sheet_names:
        if cand in names:
            return names[cand]
    return None


def write_back(table: Table, rtids: dict[int, str]) -> None:
    """rtIds in die Datei eintragen (Schluessel = Index in table.rows)."""
    if not rtids:
        return
    if table.sheet is None:
        header = list(table.header)
        rows = [list(r) for r in table.rows]
        if "rtId" not in header:
            header.insert(0, "rtId")
            rows = [[""] + r for r in rows]
        col = header.index("rtId")
        for idx, rtid in rtids.items():
            while len(rows[idx]) <= col:
                rows[idx].append("")
            rows[idx][col] = rtid
        tmp = table.path.with_suffix(table.path.suffix + ".tmp")
        with tmp.open("w", encoding="utf-8", newline="") as fh:
            w = csv.writer(fh, delimiter=table.delimiter, lineterminator="\n")
            w.writerow(header)
            w.writerows(rows)
        tmp.replace(table.path)
    else:
        openpyxl = require_openpyxl()
        wb = openpyxl.load_workbook(table.path)
        ws = wb[table.sheet]
        header = [("" if c.value is None else str(c.value).strip()) for c in ws[1]]
        if "rtId" in header:
            col = header.index("rtId") + 1
        else:
            col = len([h for h in header if h]) + 1
            ws.cell(row=1, column=col, value="rtId")
        for idx, rtid in rtids.items():
            ws.cell(row=table.first_line + idx, column=col, value=rtid)
        wb.save(table.path)


# --------------------------------------------------------------------------
# Werte umwandeln (wie das Add-in: queryController.getTableData)
# --------------------------------------------------------------------------

INT_TYPES = {"INT", "INTEGER", "INT_64", "INTEGER_64"}
DATE_TYPES = {"DATE_TIME", "DATE_TIME_OFFSET"}


def is_empty(v: Any) -> bool:
    return v is None or (isinstance(v, str) and v.strip() == "")


def upper_snake(name: str) -> str:
    return re.sub(r"(?<=[a-z0-9])(?=[A-Z])", "_", name).upper()


def enum_name(path: str, value: Any) -> str:
    """CK-Enum-Name zu einem Namen oder Schluessel; Gross-/Kleinschreibung egal."""
    names = ENUMS[path]
    if isinstance(value, (int, float)) and not isinstance(value, bool) or re.fullmatch(r"\d+", str(value)):
        k = int(float(value))
        if 0 <= k < len(names):
            return names[k]
    else:
        wanted = upper_snake(str(value).strip()).replace("_", "")
        for n in names:
            if upper_snake(n).replace("_", "") == wanted:
                return n
    allowed = ", ".join(f"{n} ({i})" for i, n in enumerate(names))
    raise ValueError(f"Spalte '{path}': '{value}' gibt es nicht. Erlaubt: {allowed}")


def convert(value: Any, vtype: str, column: str) -> Any:
    if isinstance(value, str):
        value = value.strip()
    if vtype == "STRING":
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        if isinstance(value, (dt.datetime, dt.date)):
            return value.isoformat()
        return str(value)
    if vtype == "BOOLEAN":
        if isinstance(value, bool):
            return value
        s = str(value).lower()
        if s in ("true", "1", "ja", "yes", "wahr", "x"):
            return True
        if s in ("false", "0", "nein", "no", "falsch"):
            return False
        raise ValueError(f"Spalte '{column}': '{value}' ist kein Wahrheitswert (ja/nein, true/false, 1/0)")
    if vtype == "DOUBLE":
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            return float(value)
        s = str(value).replace(" ", "")
        s = s.replace(".", "").replace(",", ".") if ("," in s and "." in s) else s.replace(",", ".")
        try:
            return float(s)
        except ValueError:
            raise ValueError(f"Spalte '{column}': '{value}' ist keine Zahl") from None
    if vtype in INT_TYPES:
        try:
            f = float(value) if isinstance(value, (int, float)) else float(str(value).replace(",", "."))
        except ValueError:
            raise ValueError(f"Spalte '{column}': '{value}' ist keine ganze Zahl") from None
        if not f.is_integer():
            raise ValueError(f"Spalte '{column}': '{value}' ist keine ganze Zahl")
        return int(f)
    if vtype == "ENUM":
        if column in ENUMS:
            return enum_name(column, value)
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            return int(value)
        return int(value) if re.fullmatch(r"\d+", str(value)) else str(value)
    if vtype in DATE_TYPES:
        return to_utc_iso(value, column)
    if vtype == "TIME_SPAN":
        return str(value)
    raise ValueError(f"Spalte '{column}': Werttyp {vtype} wird nicht unterstuetzt")


def to_utc_iso(value: Any, column: str) -> str:
    if isinstance(value, dt.datetime):
        d = value
    elif isinstance(value, dt.date):
        d = dt.datetime(value.year, value.month, value.day)
    else:
        s, d = str(value).strip(), None
        for fmt_ in ("%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M", "%d.%m.%Y"):
            try:
                d = dt.datetime.strptime(s, fmt_)
                break
            except ValueError:
                pass
        if d is None:
            try:
                d = dt.datetime.fromisoformat(s.replace("Z", "+00:00"))
            except ValueError:
                raise ValueError(f"Spalte '{column}': '{value}' ist kein Datum (TT.MM.JJJJ oder JJJJ-MM-TT)") from None
    if d.tzinfo is None:
        d = d.replace(tzinfo=LOCAL_TZ or dt.timezone.utc)
    return d.astimezone(dt.timezone.utc).strftime("%Y-%m-%dT%H:%M:%S.000Z")


def norm(v: Any) -> Any:
    """Vergleichsform fuer Datei- und Serverwert."""
    if is_empty(v):
        return None
    if isinstance(v, list):
        v = v[0] if len(v) == 1 else tuple(norm(x) for x in v)
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}T[\d:.]+(Z|[+-]\d{2}:\d{2})?", s):
        try:
            d = dt.datetime.fromisoformat(s.replace("Z", "+00:00"))
            if d.tzinfo is None:
                d = d.replace(tzinfo=dt.timezone.utc)
            return d.astimezone(dt.timezone.utc).replace(microsecond=0).isoformat()
        except ValueError:
            pass
    try:
        return float(s)
    except ValueError:
        return s.upper().replace("_", "")   # Enums: "NaturalPerson" == "NATURAL_PERSON"


def same(a: Any, b: Any) -> bool:
    return norm(a) == norm(b)


# --------------------------------------------------------------------------
# Plan: welche Zeile wird neu angelegt, aktualisiert, bleibt?
# --------------------------------------------------------------------------


@dataclass
class PlannedRow:
    index: int
    line: int
    action: str                # neu | aktualisieren | unveraendert | fehler
    rtid: str | None
    ck_type_id: str
    cells: list[dict[str, Any]]
    values: dict[str, Any]
    key: str
    changes: list[str] = field(default_factory=list)
    error: str | None = None
    matched_by_key: bool = False
    facility_rtid: str | None = None


def key_of(values: dict[str, Any], itype: ImportType) -> tuple | None:
    parts = []
    for k in itype.natural_key:
        if is_empty(values.get(k)):
            return None
        parts.append(norm(values[k]))
    return tuple(parts)


def fmt(v: Any) -> str:
    if v is None:
        return "(leer)"
    if isinstance(v, list) and len(v) == 1:
        v = v[0]
    return repr(v) if isinstance(v, str) else str(v)


def label(table: Table) -> str:
    return f"{table.path}" + (f" [{table.sheet}]" if table.sheet else "")


def check_references(row: PlannedRow, itype: ImportType, existing: Existing, refs: References) -> str | None:
    v = row.values
    if itype.key == "anlagen":
        cust = v.get(CUST_REF)
        if cust is None and row.rtid is None:
            return f"Kundennummer fehlt (Spalte {CUST_REF})"
        if cust is not None and norm(cust) not in refs.customers() and norm(cust) not in refs.pending_customers:
            return f"Kunde {cust} gibt es nicht. Erst die Kunden importieren."
        # Jede Anlage braucht einen Elternknoten (ParentChild, Multiplizitaet One) - im Labor
        # den Ort aus Locations.Austria, gefunden ueber die PLZ.
        zipcode = v.get(CITY_REF)
        if zipcode is None:
            if row.rtid is None:
                return f"PLZ fuer den Ort fehlt (Spalte {CITY_REF}). Jede Anlage braucht einen Ort."
        else:
            zips = refs.zipcodes()
            n = zips.get(norm(zipcode), 0)
            if not zips:
                return ("Im Tenant gibt es kein Ortsverzeichnis. Blueprint Locations.Austria installieren "
                        "(macht scripts/neue-energiegemeinschaft.sh).")
            if n == 0:
                return f"PLZ {zipcode} steht nicht im Ortsverzeichnis (Locations.Austria)."
            if n > 1:
                return (f"PLZ {zipcode} gehoert zu {n} Orten, die Zuordnung ist nicht eindeutig. "
                        f"Diese Anlage in der Oberflaeche anlegen (Ort dort auswaehlen).")
    if itype.entity:
        name, cust = v.get(FAC_NAME), v.get(FAC_CUST)
        if name is None and cust is None:
            if row.rtid is None:
                return "Anlage fehlt (Spalten Anlagenname und Kundennummer)."
            return None
        if name is None or cust is None:
            return "Fuer die Anlage braucht es Anlagenname UND Kundennummer."
        k = (norm(name), norm(cust))
        hits = refs.facilities().get(k, [])
        if len(hits) == 1:
            row.facility_rtid = hits[0]
        elif len(hits) > 1:
            return f"Anlage '{name}' von Kunde {cust} gibt es {len(hits)}-mal. Bitte in der Oberflaeche bereinigen."
        elif k in refs.pending_facilities:
            row.facility_rtid = "(wird angelegt)"
        else:
            return f"Anlage '{name}' von Kunde {cust} gibt es nicht. Erst die Anlagen importieren."
    return None


def plan(table: Table, existing: Existing, itype: ImportType, use_key_match: bool,
         refs: References) -> list[PlannedRow]:
    unknown = [h for h in table.header
               if h and h not in existing.columns and h not in SPECIAL_COLUMNS and not h.startswith(("#", "_"))]
    if unknown:
        fail(f"{label(table)}: diese Spalten kennt der Typ {itype.key} nicht: {', '.join(unknown)}\n"
             f"  Erlaubt: rtId, {', '.join(existing.column_order)}\n"
             f"  (Spalten, die mit # oder _ beginnen, werden ignoriert)")
    missing_key = [k for k in itype.natural_key if k not in table.header]
    if missing_key and use_key_match:
        print(f"   Hinweis: Schluesselspalte(n) {', '.join(missing_key)} fehlen, kein Abgleich ueber Schluessel.")

    data_cols = [(i, h) for i, h in enumerate(table.header) if h in existing.columns]
    rtid_col = table.header.index("rtId") if "rtId" in table.header else None
    ck_col = table.header.index("ckTypeId") if "ckTypeId" in table.header else None

    server_by_key: dict[tuple, str] = {}
    for rtid, cells in existing.rows.items():
        k = key_of(cells, itype)
        if k is not None:
            server_by_key.setdefault(k, rtid)

    planned: list[PlannedRow] = []
    seen_keys: dict[tuple, int] = {}
    for idx, raw in enumerate(table.rows):
        line = table.first_line + idx
        raw = list(raw) + [None] * (len(table.header) - len(raw))
        if all(is_empty(raw[i]) for i, _ in data_cols):
            continue
        rtid = None if rtid_col is None or is_empty(raw[rtid_col]) else str(raw[rtid_col]).strip()
        ck_type = existing.ck_type_id
        if ck_col is not None and not is_empty(raw[ck_col]) and not itype.entity:
            ck_type = str(raw[ck_col]).strip()

        cells: list[dict[str, Any]] = []
        values: dict[str, Any] = {}
        error = None
        for i, col in data_cols:
            v = raw[i]
            if is_empty(v):
                if "->" in col:
                    continue      # leerer Verweis: nicht schicken (sonst MAPPING001)
                cells.append({"attributePath": col, "value": None})
                values[col] = None
                continue
            try:
                cv = convert(v, existing.columns[col], col)
            except ValueError as e:
                error = str(e)
                break
            cells.append({"attributePath": col, "value": cv})
            values[col] = cv

        k = key_of(values, itype) if not error else None
        raw_by_col = {h: raw[i] for i, h in enumerate(table.header)}
        key_parts = [values.get(c) if c in values else raw_by_col.get(c) for c in itype.natural_key]
        key_txt = " / ".join("?" if is_empty(x) else str(x) for x in key_parts) \
            if any(not is_empty(x) for x in key_parts) else "(ohne Schluessel)"
        row = PlannedRow(index=idx, line=line, action="neu", rtid=rtid, ck_type_id=ck_type,
                         cells=cells, values=values, key=key_txt)
        if error:
            row.action, row.error = "fehler", error
            planned.append(row)
            continue
        if k is not None:
            if k in seen_keys:
                row.action = "fehler"
                row.error = f"Schluessel {key_txt} steht doppelt in der Datei (auch Zeile {seen_keys[k]})"
                planned.append(row)
                continue
            seen_keys[k] = line

        if rtid is None and use_key_match and k is not None and k in server_by_key:
            rtid = server_by_key[k]
            row.rtid, row.matched_by_key = rtid, True

        if rtid is not None:
            server = existing.rows.get(rtid)
            if server is None:
                row.action = "fehler"
                row.error = (f"rtId {rtid} gibt es in diesem Tenant nicht (aus einem anderen Tenant kopiert? "
                             f"Dann die rtId-Spalte leeren)")
                planned.append(row)
                continue
            for c in cells:
                old = server.get(c["attributePath"])
                if not same(old, c["value"]):
                    row.changes.append(f"{c['attributePath']}: {fmt(old)} -> {fmt(c['value'])}")
            row.action = "aktualisieren" if row.changes else "unveraendert"

        if row.action in ("neu", "aktualisieren"):
            ref_error = check_references(row, itype, existing, refs)
            if ref_error:
                row.action, row.error = "fehler", ref_error
        planned.append(row)
    return planned


# --------------------------------------------------------------------------
# Schreiben
# --------------------------------------------------------------------------


def write_query_rows(conn: Connection, existing: Existing, creates: list[PlannedRow],
                     updates: list[PlannedRow]) -> None:
    for start in range(0, len(creates), BATCH_SIZE):
        batch = creates[start:start + BATCH_SIZE]
        entities = [{"ckTypeId": r.ck_type_id, "cells": r.cells} for r in batch]
        try:
            data = graphql(conn, M_CREATE, {"queryRtId": existing.query_rtid, "entities": entities})
        except GraphQlError as e:
            fail(f"{e}\n   Betroffen: Zeilen {lines_of(batch)}. Keine dieser Zeilen wurde angelegt.")
        created = ((data.get("runtime") or {}).get("runtimeQuery") or {}).get("create") or []
        if len(created) != len(batch):
            fail(f"Server meldet {len(created)} angelegte Objekte, erwartet {len(batch)}.")
        fixups = []
        for r, c in zip(batch, created):
            r.rtid = c["rtId"]
            # Zaehler (System/AutoIncrement) ersetzen Werte beim Anlegen, z. B. die Kundennummer.
            # Der Dateiwert gilt: sofort nachtragen, sonst finden Anlagen ihren Kunden nicht.
            got = {x["attributePath"]: x.get("value") for x in ((c.get("cells") or {}).get("items") or [])}
            for cell in r.cells:
                path = cell["attributePath"]
                if "->" in path or cell["value"] is None or path not in got:
                    continue
                if not same(got[path], cell["value"]):
                    fixups.append(r)
                    print(f"   Zeile {r.line:>4}: {path} vom Server als {fmt(got[path])} vergeben, "
                          f"setze Dateiwert {fmt(cell['value'])}")
                    break
        if fixups:
            write_query_updates(conn, existing, fixups)
    write_query_updates(conn, existing, updates)


def write_query_updates(conn: Connection, existing: Existing, rows: list[PlannedRow]) -> None:
    for start in range(0, len(rows), BATCH_SIZE):
        batch = rows[start:start + BATCH_SIZE]
        entities = [{"rtId": r.rtid, "item": {"ckTypeId": r.ck_type_id, "cells": r.cells}} for r in batch]
        try:
            graphql(conn, M_UPDATE, {"queryRtId": existing.query_rtid, "entities": entities})
        except GraphQlError as e:
            fail(f"{e}\n   Betroffen: Zeilen {lines_of(batch)}. Keine dieser Zeilen wurde aktualisiert.")


def entity_input(spec: EntitySpec, row: PlannedRow, existing: Existing, create: bool) -> dict[str, Any]:
    item: dict[str, Any] = {}
    for cell in row.cells:
        path, value = cell["attributePath"], cell["value"]
        if "->" in path:
            continue
        if path in ENUMS and value is not None:
            value = upper_snake(value)          # GraphQL-Enum: Solar -> SOLAR
        item[path] = value
    if create and item.get("carrierType") is None:
        item["carrierType"] = "ELECTRICITY"     # Pflichtfeld ohne Vorgabewert im CK
    old_fac = existing.facility_of.get(row.rtid or "")
    if row.facility_rtid and row.facility_rtid != old_fac:
        assoc = []
        if old_fac:
            assoc.append({"target": {"rtId": old_fac, "ckTypeId": FAC_CK}, "modOption": "DELETE"})
        assoc.append({"target": {"rtId": row.facility_rtid, "ckTypeId": FAC_CK}, "modOption": "CREATE"})
        item["parent"] = assoc
    return item


def write_entities(conn: Connection, spec: EntitySpec, existing: Existing, creates: list[PlannedRow],
                   updates: list[PlannedRow]) -> None:
    for start in range(0, len(creates), BATCH_SIZE):
        batch = creates[start:start + BATCH_SIZE]
        entities = [entity_input(spec, r, existing, True) for r in batch]
        try:
            data = graphql(conn, m_entities(spec, "create"), {"entities": entities})
        except GraphQlError as e:
            fail(f"{e}\n   Betroffen: Zeilen {lines_of(batch)}. Keine dieser Zeilen wurde angelegt.")
        created = ((data.get("runtime") or {}).get(spec.mutation_field) or {}).get("create") or []
        by_number = {norm(c.get("meteringPointNumber")): c["rtId"] for c in created}
        for r in batch:
            r.rtid = by_number.get(norm(r.values.get("meteringPointNumber")))
            if not r.rtid:
                fail(f"Zeile {r.line}: Server meldet keinen angelegten Zaehlpunkt {r.key}.")
    for start in range(0, len(updates), BATCH_SIZE):
        batch = updates[start:start + BATCH_SIZE]
        entities = [{"rtId": r.rtid, "item": entity_input(spec, r, existing, False)} for r in batch]
        try:
            graphql(conn, m_entities(spec, "update"), {"entities": entities})
        except GraphQlError as e:
            fail(f"{e}\n   Betroffen: Zeilen {lines_of(batch)}. Keine dieser Zeilen wurde aktualisiert.")


def lines_of(batch: list[PlannedRow]) -> str:
    return ", ".join(str(r.line) for r in batch[:20]) + (" ..." if len(batch) > 20 else "")


def run_import(conn: Connection, itype: ImportType, table: Table, args: argparse.Namespace,
               refs: References) -> dict[str, int]:
    print(f"\n== {itype.label}: {label(table)}")
    existing = load_existing(conn, itype)
    via = f"Abfrage {itype.query_name} ({existing.query_rtid})" if not itype.entity \
        else f"Entity-API {itype.entity.mutation_field}"
    print(f"   {via}, Typ {existing.ck_type_id}, {len(existing.rows)} Objekte im Tenant {conn.tenant}")
    rows = plan(table, existing, itype, not args.ohne_schluesselabgleich, refs)
    stats = {"neu": 0, "aktualisieren": 0, "unveraendert": 0, "fehler": 0}
    for r in rows:
        stats[r.action] += 1
        if r.action == "unveraendert" and not args.ausfuehrlich:
            continue
        tag = {"neu": "NEU", "aktualisieren": "AKTUALISIEREN", "unveraendert": "unveraendert",
               "fehler": "FEHLER"}[r.action]
        extra = " (vorhanden, ueber Schluessel gefunden)" if r.matched_by_key and r.action != "fehler" else ""
        print(f"   Zeile {r.line:>4}: {tag:<13} {r.key}{extra}")
        if r.error:
            print(f"              {r.error}")
        for ch in r.changes:
            print(f"              {ch}")
    print(f"   Summe: {stats['neu']} neu, {stats['aktualisieren']} zu aktualisieren, "
          f"{stats['unveraendert']} unveraendert, {stats['fehler']} Fehler")

    if itype.key == "kunden":
        refs.pending_customers |= {norm(r.values.get("customerNumber")) for r in rows if r.action == "neu"}
    if itype.key == "anlagen":
        refs.pending_facilities |= {(norm(r.values.get("name")), norm(r.values.get(CUST_REF)))
                                    for r in rows if r.action == "neu"}

    if stats["fehler"]:
        print("   Abbruch fuer diese Datei: erst die Fehler beheben. Es wurde nichts geschrieben.")
        return stats
    if args.probelauf:
        print("   Probelauf: nichts geschrieben.")
        return stats

    creates = [r for r in rows if r.action == "neu"]
    updates = [r for r in rows if r.action == "aktualisieren"]
    if itype.entity:
        write_entities(conn, itype.entity, existing, creates, updates)
    else:
        write_query_rows(conn, existing, creates, updates)
    refs.invalidate()
    print(f"   Geschrieben: {len(creates)} angelegt, {len(updates)} aktualisiert.")

    new_rtids = {r.index: r.rtid for r in rows
                 if r.rtid and (r.action == "neu" or r.matched_by_key)}
    if new_rtids and not args.nicht_zurueckschreiben:
        write_back(table, new_rtids)
        print(f"   rtIds in {label(table)} eingetragen ({len(new_rtids)} Zeilen).")
    return stats


def tables_for(target: str, path: Path) -> list[tuple[ImportType, Table]]:
    if not path.exists():
        fail(f"{path} gibt es nicht.")
    is_xlsx = path.suffix.lower() in (".xlsx", ".xlsm")
    if target != "alle":
        itype = TYPES_BY_KEY[target]
        if path.is_dir():
            fail(f"{path} ist ein Ordner. Fuer einen Ordner 'alle' verwenden.")
        if is_xlsx:
            sheet = xlsx_sheet_for(path, itype) or require_openpyxl().load_workbook(path, read_only=True).sheetnames[0]
            return [(itype, read_xlsx(path, sheet))]
        return [(itype, read_csv(path))]

    result: list[tuple[ImportType, Table]] = []
    if path.is_dir():
        for itype in TYPES:
            csv_file, xlsx_file = path / f"{itype.key}.csv", path / f"{itype.key}.xlsx"
            if csv_file.exists():
                result.append((itype, read_csv(csv_file)))
            elif xlsx_file.exists():
                sheet = require_openpyxl().load_workbook(xlsx_file, read_only=True).sheetnames[0]
                result.append((itype, read_xlsx(xlsx_file, sheet)))
        if not result:
            fail(f"In {path} liegt keine der Dateien {', '.join(t.key + '.csv' for t in TYPES)}.")
        return result
    if is_xlsx:
        for itype in TYPES:
            sheet = xlsx_sheet_for(path, itype)
            if sheet:
                result.append((itype, read_xlsx(path, sheet)))
        if not result:
            fail(f"{path}: keine Blaetter namens Kunden / Anlagen / Verbraucher / Erzeuger gefunden.")
        return result
    fail("'alle' braucht einen Ordner mit kunden.csv, anlagen.csv, ... oder eine .xlsx-Mappe.")
    return []


def cmd_columns(conn: Connection, target: str) -> None:
    for itype in (TYPES if target == "alle" else [TYPES_BY_KEY[target]]):
        existing = load_existing(conn, itype)
        print(f"{itype.key} ({existing.ck_type_id}, {len(existing.rows)} vorhanden):")
        for c in existing.column_order:
            hint = f"  Werte: {', '.join(ENUMS[c])}" if c in ENUMS else ""
            print(f"  {c:<88} {existing.columns[c]}{hint}")
        print("  Kopfzeile: rtId;" + ";".join(existing.column_order))


def main(argv: list[str]) -> int:
    p = argparse.ArgumentParser(
        prog="import-stammdaten.py",
        description="Stammdaten (Kunden, Anlagen, Zaehlpunkte) in eine Energiegemeinschaft importieren.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="Beispiele:" + __doc__.split("Beispiele:")[1])
    p.add_argument("-c", "--context", required=True, help="octo-cli-Kontext des Gemeinschafts-Tenants")
    p.add_argument("aktion", choices=["kunden", "anlagen", "verbraucher", "erzeuger", "alle", "spalten"],
                   help="Typ der Datei; 'alle' = Ordner oder Mappe; 'spalten' = gueltige Spalten anzeigen")
    p.add_argument("datei", nargs="?", help="CSV/XLSX-Datei bzw. Ordner (bei 'spalten': Typ oder leer)")
    p.add_argument("-n", "--probelauf", action="store_true", help="nur anzeigen, was passieren wuerde")
    p.add_argument("--ohne-schluesselabgleich", action="store_true",
                   help="Zeilen ohne rtId immer neu anlegen (Verhalten des Add-ins)")
    p.add_argument("--nicht-zurueckschreiben", action="store_true", help="rtIds nicht in die Datei schreiben")
    p.add_argument("-v", "--ausfuehrlich", action="store_true", help="auch unveraenderte Zeilen zeigen")
    p.add_argument("--assets-url", help="Asset-Dienst ueberschreiben, z. B. https://assets.<domain>/")
    p.add_argument("--tenant", help="Tenant ueberschreiben (Standard: Tenant des Kontexts)")
    p.add_argument("--unsicher", action="store_true", help="TLS-Zertifikat nicht pruefen (nur lokal!)")
    p.add_argument("--kein-refresh", action="store_true", help=argparse.SUPPRESS)
    args = p.parse_args(argv)
    try:
        sys.stdout.reconfigure(line_buffering=True)
    except AttributeError:
        pass

    try:
        if args.aktion == "spalten":
            target = (args.datei or "alle").lower()
            if target != "alle" and target not in TYPES_BY_KEY:
                fail(f"Unbekannter Typ '{target}'. Moeglich: {', '.join(TYPES_BY_KEY)}, alle")
            cmd_columns(load_connection(args), target)
            return 0
        if not args.datei:
            fail("Datei bzw. Ordner fehlt.")
        tables = tables_for(args.aktion, Path(args.datei))
        conn = load_connection(args)
        refs = References(conn)
        print(f"Ziel: Tenant {conn.tenant} ({conn.graphql_url})" + ("  [PROBELAUF]" if args.probelauf else ""))
        total = {"neu": 0, "aktualisieren": 0, "unveraendert": 0, "fehler": 0}
        for itype, table in tables:
            stats = run_import(conn, itype, table, args, refs)
            for k in total:
                total[k] += stats[k]
            if stats["fehler"]:
                if len(tables) > 1:
                    print("\nAbbruch: die folgenden Dateien haengen von dieser ab.")
                return 2
        if len(tables) > 1:
            print(f"\nGesamt: {total['neu']} neu, {total['aktualisieren']} zu aktualisieren/aktualisiert, "
                  f"{total['unveraendert']} unveraendert.")
        return 0
    except UserError as e:
        print(f"\nFEHLER: {e}", file=sys.stderr)
        return 1
    except KeyboardInterrupt:
        return 130


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
